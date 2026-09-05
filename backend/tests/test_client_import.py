"""Import d'historique commercial — le fichier, ses refus, et son effet réel.

Le test le plus important de ce fichier n'est pas un test d'import : c'est
`test_l_historique_verse_atteint_le_moteur_de_cycles`. La règle du projet est
qu'une donnée saisissable doit être vérifiée jusqu'à son effet — écrire des
lignes en base sans que le moteur les lise donnerait un import parfaitement
fonctionnel et parfaitement inutile, du même genre que la courbe de dividende
qu'on pouvait saisir sans qu'elle change le prix.

Le reste tient sur trois refus : rien ne s'écrit avant d'avoir été montré,
rien ne se duplique au second passage, et une transaction se rattache à
l'employeur de SA date, pas à celui d'aujourd'hui.
"""
import io
import json
from datetime import date

import pytest
from fastapi.testclient import TestClient
from sqlmodel import Session, SQLModel, create_engine, select

from backend.app.api.auth import get_current_user
from backend.app.core.client_cycle import compute_cycle
from backend.app.core.client_intelligence import (
    _dates_de_trade, deals_of_affiliations, deals_of_client, observed_behaviour,
)
from backend.app.db.database import get_session
from backend.app.db.models import (
    Affiliation, Client, ClientImportBatch, ClientMandate, ClientTradeHistory,
    Person, User,
)


@pytest.fixture
def app_client(monkeypatch):
    import backend.app.db.database as db
    from sqlalchemy.pool import StaticPool

    engine = create_engine("sqlite://", poolclass=StaticPool,
                           connect_args={"check_same_thread": False})
    SQLModel.metadata.create_all(engine)
    monkeypatch.setattr(db, "engine", engine)

    from backend.app.main import app

    session = Session(engine)
    session.add(User(id=1, username="alice", email="a@d.com", password_hash="x",
                     role="user", entity_id=1))
    session.commit()

    app.dependency_overrides[get_session] = lambda: session
    app.dependency_overrides[get_current_user] = lambda: session.get(User, 1)
    client = TestClient(app)
    client.session = session
    yield client
    app.dependency_overrides.clear()
    session.close()


def _envoyer(app_client, charge: dict, *, route="preview", **form):
    contenu = json.dumps(charge, ensure_ascii=False).encode("utf-8")
    return app_client.post(
        f"/api/client-import/{route}",
        files={"file": ("historique.json", io.BytesIO(contenu),
                        "application/json")},
        data=form)


DOSSIER_COMPLET = {
    "clients": [{"name": "Bank A", "client_type": "private_bank",
                 "country": "Suisse", "status": "active"},
                {"name": "Bank B", "client_type": "asset_manager"}],
    "contacts": [{"first_name": "Jean", "last_name": "Dupont",
                  "email": "jean.dupont@bank-a.com"}],
    "affiliations": [
        {"person_email_or_name": "jean.dupont@bank-a.com",
         "client_name": "Bank A", "job_title": "Gérant",
         "commercial_role": "portfolio_manager",
         "start_date": "2022-01-01", "end_date": "2025-12-31"},
        {"person_email_or_name": "jean.dupont@bank-a.com",
         "client_name": "Bank B", "job_title": "CIO",
         "commercial_role": "cio", "start_date": "2026-01-01"},
    ],
    "transactions": [
        {"client_name": "Bank A", "person_email_or_name": "jean.dupont@bank-a.com",
         "trade_date": "2024-01-15", "maturity_date": "2027-01-15",
         "product_type": "Autocall Athena", "underlying": "SX5E",
         "issuer": "BNP Paribas", "currency": "EUR", "notional": 1000000,
         "external_ref": "XS001"},
        {"client_name": "Bank A", "person_email_or_name": "jean.dupont@bank-a.com",
         "trade_date": "2024-04-15", "product_type": "Autocall Athena",
         "notional": 1000000, "external_ref": "XS002"},
        {"client_name": "Bank A", "person_email_or_name": "jean.dupont@bank-a.com",
         "trade_date": "2024-07-15", "product_type": "Phoenix Memory",
         "notional": 2000000, "external_ref": "XS003"},
        {"client_name": "Bank A", "person_email_or_name": "jean.dupont@bank-a.com",
         "trade_date": "2024-10-14", "product_type": "Autocall Athena",
         "notional": 1000000, "external_ref": "XS004"},
    ],
}


# ── Rien ne s'écrit avant d'avoir été montré ─────────────────────────

def test_l_apercu_annonce_sans_rien_ecrire(app_client):
    reponse = _envoyer(app_client, DOSSIER_COMPLET)
    assert reponse.status_code == 200, reponse.text
    corps = reponse.json()

    assert corps["created"]["clients"] == 2
    assert corps["created"]["contacts"] == 1
    assert corps["created"]["affiliations"] == 2
    assert corps["created"]["transactions"] == 4
    assert corps["blocking"] is False

    # Et la base n'a pas bougé.
    assert app_client.session.exec(select(Client)).all() == []
    assert app_client.session.exec(select(Person)).all() == []


def test_le_versement_ecrit_ce_que_l_apercu_avait_annonce(app_client):
    apercu = _envoyer(app_client, DOSSIER_COMPLET).json()
    verse = _envoyer(app_client, DOSSIER_COMPLET, route="apply").json()

    assert verse["created"] == apercu["created"]
    assert len(app_client.session.exec(select(Client)).all()) == 2
    assert len(app_client.session.exec(select(Affiliation)).all()) == 2
    assert len(app_client.session.exec(select(ClientTradeHistory)).all()) == 4


def test_l_import_conserve_le_mandat_le_cadre_juridique_et_le_payoff(app_client):
    charge = {
        "clients": [{"name": "Client Institutionnel", "status": "active"}],
        "mandates": [{
            "client_name": "Client Institutionnel", "name": "Mandat Taux",
            "mandate_type": "mandate", "reference_currency": "EUR",
        }],
        "transactions": [{
            "client_name": "Client Institutionnel", "mandate_name": "Mandat Taux",
            "trade_date": "2026-04-20", "transaction_format": "OTC",
            "instrument_family": "Swap", "payoff_family": "Swap",
            "payoff_description": "Swap CMS 10Y / 2Y",
            "documentation_reference": "ISDA-2025-CLIENT",
            "product_type": "CMS steepener", "external_ref": "LEGACY-OTC-1",
        }],
    }

    response = _envoyer(app_client, charge, route="apply")
    assert response.status_code == 200, response.text
    mandate = app_client.session.exec(select(ClientMandate)).one()
    trade = app_client.session.exec(select(ClientTradeHistory)).one()
    assert trade.mandate_id == mandate.id
    assert (trade.transaction_format, trade.instrument_family,
            trade.payoff_family) == ("OTC", "Swap", "Swap")
    assert trade.documentation_reference == "ISDA-2025-CLIENT"
    assert trade.payoff_description == "Swap CMS 10Y / 2Y"


# ── Le test qui compte : l'effet réel ────────────────────────────────

def test_l_historique_verse_atteint_le_moteur_de_cycles(app_client):
    """Écrire des lignes que le moteur ne lit pas donnerait un import
    parfaitement fonctionnel et parfaitement inutile."""
    session = app_client.session
    client_avant = session.exec(select(Client)).all()
    assert client_avant == []

    _envoyer(app_client, DOSSIER_COMPLET, route="apply")

    bank_a = session.exec(select(Client).where(Client.name == "Bank A")).first()
    transactions = deals_of_client(session, bank_a.id)
    assert len(transactions) == 4

    cycle = compute_cycle(_dates_de_trade(transactions), asof=date(2025, 1, 1))
    # Quatre transactions à trois mois d'intervalle : une cadence trimestrielle,
    # là où sans import le moteur aurait dit « historique insuffisant ».
    assert cycle.has_history
    assert cycle.cadence == "trimestrielle"
    assert round(cycle.median_interval_days) == 91

    comportement = observed_behaviour(transactions)
    assert comportement["n_trades"] == 4
    assert comportement["n_imported"] == 4
    assert comportement["product_types"][0] == ("Autocall Athena", 3)


def test_l_ecran_sait_distinguer_l_importe_du_booke(app_client):
    """Une lecture fondée surtout sur de l'historique versé n'a pas la même
    valeur de preuve qu'une lecture fondée sur nos propres trades."""
    _envoyer(app_client, DOSSIER_COMPLET, route="apply")
    session = app_client.session
    bank_a = session.exec(select(Client).where(Client.name == "Bank A")).first()
    lecture = observed_behaviour(deals_of_client(session, bank_a.id))
    assert lecture["n_imported"] == lecture["n_trades"]
    assert "historique versé" in " ".join(lecture["explanation"])


# ── L'employeur de la DATE, pas celui d'aujourd'hui ──────────────────

def test_une_transaction_se_rattache_a_l_affiliation_de_sa_date(app_client):
    """Jean est chez Bank B aujourd'hui ; ses trades de 2024 doivent pointer sur
    son passage chez Bank A, sinon tout l'intérêt du modèle tombe."""
    _envoyer(app_client, DOSSIER_COMPLET, route="apply")
    session = app_client.session

    bank_a = session.exec(select(Client).where(Client.name == "Bank A")).first()
    chez_a = session.exec(
        select(Affiliation).where(Affiliation.client_id == bank_a.id)).first()

    lignes = session.exec(select(ClientTradeHistory)).all()
    assert {l.affiliation_id for l in lignes} == {chez_a.id}
    assert len(deals_of_affiliations(session, [chez_a.id])) == 4


def test_une_transaction_hors_de_tout_passage_est_refusee(app_client):
    """Une transaction datée d'une période où la personne ne travaillait pas
    là est une erreur de saisie, pas une donnée à ranger n'importe où."""
    charge = {
        "clients": [{"name": "Bank A"}],
        "contacts": [{"first_name": "Jean", "last_name": "Dupont",
                      "email": "j@a.com"}],
        "affiliations": [{"person_email_or_name": "j@a.com",
                          "client_name": "Bank A", "start_date": "2026-01-01"}],
        "transactions": [{"client_name": "Bank A", "person_email_or_name": "j@a.com",
                          "trade_date": "2020-05-05"}],
    }
    corps = _envoyer(app_client, charge).json()
    assert corps["blocking"] is True
    anomalie = corps["issues"][0]
    assert anomalie["feuille"] == "transactions"
    assert "ne couvre le 2020-05-05" in anomalie["message"]


# ── Rapprocher plutôt que dupliquer ──────────────────────────────────

def test_verser_deux_fois_le_meme_fichier_ne_double_rien(app_client):
    """Le pire résultat possible pour un outil censé consolider."""
    _envoyer(app_client, DOSSIER_COMPLET, route="apply")
    second = _envoyer(app_client, DOSSIER_COMPLET, route="apply").json()

    session = app_client.session
    assert len(session.exec(select(Client)).all()) == 2
    assert len(session.exec(select(Person)).all()) == 1
    assert len(session.exec(select(Affiliation)).all()) == 2
    assert len(session.exec(select(ClientTradeHistory)).all()) == 4

    assert second["total_created"] == 0
    assert second["updated"]["clients"] == 2
    assert second["skipped"]["affiliations"] == 2
    assert second["skipped"]["transactions"] == 4


def test_un_import_complete_les_champs_vides_sans_ecraser_la_saisie(app_client):
    """Le fichier est une source d'appoint, pas la vérité qui prime."""
    session = app_client.session
    session.add(Client(name="Bank A", entity_id=1, country="France",
                       notes="Note saisie à la main"))
    session.commit()

    _envoyer(app_client, {"clients": [
        {"name": "Bank A", "country": "Suisse", "legal_name": "Bank A SA"}]},
        route="apply")

    session.expire_all()
    fiche = session.exec(select(Client).where(Client.name == "Bank A")).first()
    assert fiche.country == "France"           # la saisie manuelle tient
    assert fiche.legal_name == "Bank A SA"     # le vide est complété
    assert fiche.notes == "Note saisie à la main"


# ── Les refus de format ──────────────────────────────────────────────

def test_une_ligne_fautive_bloque_tout_le_versement(app_client):
    """Un import à moitié passé laisse une base dans un état que personne ne
    connaît."""
    charge = {"clients": [{"name": "Bank A"}, {"name": ""}]}
    reponse = _envoyer(app_client, charge, route="apply")
    assert reponse.status_code == 422
    assert "rien n'a été importé" in reponse.json()["detail"]
    assert app_client.session.exec(select(Client)).all() == []


def test_on_peut_demander_explicitement_d_ignorer_les_lignes_fautives(app_client):
    charge = {"clients": [{"name": "Bank A"}, {"name": ""}]}
    reponse = _envoyer(app_client, charge, route="apply", skip_invalid="true")
    assert reponse.status_code == 200
    assert len(app_client.session.exec(select(Client)).all()) == 1


def test_un_vocabulaire_inconnu_est_refuse_avec_les_valeurs_acceptees(app_client):
    corps = _envoyer(app_client, {
        "clients": [{"name": "X", "client_type": "hedge_fund"}]}).json()
    assert corps["blocking"] is True
    assert "asset_manager" in corps["issues"][0]["message"]


def test_une_date_illisible_est_signalee_avec_son_numero_de_ligne(app_client):
    corps = _envoyer(app_client, {
        "clients": [{"name": "Bank A"}],
        "interactions": [
            {"client_name": "Bank A", "interaction_date": "2026-01-01"},
            {"client_name": "Bank A", "interaction_date": "hier"},
        ]}).json()
    anomalies = [a for a in corps["issues"] if a["feuille"] == "interactions"]
    assert len(anomalies) == 1
    assert anomalies[0]["ligne"] == 2
    assert "hier" in anomalies[0]["message"]


def test_les_dates_francaises_sont_acceptees(app_client):
    """Un export d'un autre outil ne sortira pas forcément en ISO."""
    corps = _envoyer(app_client, {
        "clients": [{"name": "Bank A"}],
        "interactions": [{"client_name": "Bank A",
                          "interaction_date": "15/03/2024"}]}).json()
    assert corps["blocking"] is False
    assert corps["created"]["interactions"] == 1


def test_les_entetes_tolerent_la_casse_et_les_accents(app_client):
    """L'export d'un autre outil ne respectera jamais notre casse."""
    corps = _envoyer(app_client, {"CLIENTS": [{"Name": "Bank A",
                                               "Client Type": "bank"}]}).json()
    assert corps["created"]["clients"] == 1


def test_une_section_inconnue_est_signalee_sans_bloquer(app_client):
    corps = _envoyer(app_client, {"clients": [{"name": "X"}],
                                  "portefeuilles": [{"x": 1}]}).json()
    assert corps["unknown_sheets"] == ["portefeuilles"]
    assert corps["blocking"] is False


def test_un_fichier_sans_section_exploitable_est_refuse(app_client):
    reponse = _envoyer(app_client, {"n_importe_quoi": []})
    assert reponse.status_code == 422
    assert "Aucune section exploitable" in reponse.json()["detail"]


def test_un_json_illisible_est_refuse_clairement(app_client):
    reponse = app_client.post(
        "/api/client-import/preview",
        files={"file": ("x.json", io.BytesIO(b"{pas du json"), "application/json")})
    assert reponse.status_code == 422
    assert "JSON illisible" in reponse.json()["detail"]


def test_une_extension_inconnue_est_refusee(app_client):
    reponse = app_client.post(
        "/api/client-import/preview",
        files={"file": ("historique.csv", io.BytesIO(b"a,b"), "text/csv")})
    assert reponse.status_code == 422
    assert ".xlsx" in reponse.json()["detail"]


# ── Le modèle à remplir ──────────────────────────────────────────────

def test_le_modele_excel_porte_une_feuille_par_objet(app_client):
    from openpyxl import load_workbook
    reponse = app_client.get("/api/client-import/template.xlsx")
    assert reponse.status_code == 200
    classeur = load_workbook(io.BytesIO(reponse.content))
    for feuille in ("clients", "contacts", "affiliations", "transactions",
                    "interactions"):
        assert feuille in classeur.sheetnames
    assert "Mode d'emploi" in classeur.sheetnames


def test_le_modele_excel_se_reimporte_tel_quel(app_client):
    """Le meilleur test d'un gabarit : qu'il passe l'import qu'il prépare.
    Une colonne mal nommée dans le modèle se verrait immédiatement."""
    modele = app_client.get("/api/client-import/template.xlsx").content
    reponse = app_client.post(
        "/api/client-import/preview",
        files={"file": ("modele.xlsx", io.BytesIO(modele),
                        "application/vnd.openxmlformats-officedocument"
                        ".spreadsheetml.sheet")})
    assert reponse.status_code == 200, reponse.text
    corps = reponse.json()
    assert corps["blocking"] is False, corps["issues"]
    assert corps["created"]["clients"] == 1
    assert corps["created"]["affiliations"] == 1
    assert corps["created"]["transactions"] == 1


def test_le_modele_json_porte_ses_vocabulaires(app_client):
    corps = app_client.get("/api/client-import/template.json").json()
    assert "asset_manager" in corps["_format"]["vocabulaires"]["client_type"]
    assert corps["_format"]["champs_requis"]["transactions"] == [
        "client_name", "trade_date"]


# ── Traçabilité et annulation ────────────────────────────────────────

def test_un_versement_laisse_un_lot_consultable(app_client):
    _envoyer(app_client, DOSSIER_COMPLET, route="apply")
    lots = app_client.get("/api/client-import/batches").json()
    assert len(lots) == 1
    assert lots[0]["status"] == "applied"
    assert lots[0]["rows_created"] == 9
    assert lots[0]["report"]["created"]["transactions"] == 4


def test_annuler_un_lot_retire_ses_transactions_mais_garde_le_referentiel(app_client):
    """Sociétés et personnes ont pu recevoir depuis des saisies manuelles ;
    les supprimer emporterait ce travail."""
    verse = _envoyer(app_client, DOSSIER_COMPLET, route="apply").json()
    reponse = app_client.post(
        f"/api/client-import/batches/{verse['batch_id']}/revert")
    assert reponse.status_code == 200
    assert reponse.json()["trades_removed"] == 4

    session = app_client.session
    assert session.exec(select(ClientTradeHistory)).all() == []
    assert len(session.exec(select(Client)).all()) == 2
    assert len(session.exec(select(Affiliation)).all()) == 2


def test_un_lot_deja_annule_ne_se_rannule_pas(app_client):
    verse = _envoyer(app_client, DOSSIER_COMPLET, route="apply").json()
    app_client.post(f"/api/client-import/batches/{verse['batch_id']}/revert")
    seconde = app_client.post(
        f"/api/client-import/batches/{verse['batch_id']}/revert")
    assert seconde.status_code == 409


def test_apres_annulation_le_moteur_redevient_muet(app_client):
    """La vérification symétrique de l'effet : si l'import fait parler le
    moteur, l'annulation doit le faire taire."""
    verse = _envoyer(app_client, DOSSIER_COMPLET, route="apply").json()
    session = app_client.session
    bank_a = session.exec(select(Client).where(Client.name == "Bank A")).first()
    assert compute_cycle(_dates_de_trade(deals_of_client(session, bank_a.id)),
                         asof=date(2025, 1, 1)).has_history

    app_client.post(f"/api/client-import/batches/{verse['batch_id']}/revert")
    session.expire_all()
    apres = compute_cycle(_dates_de_trade(deals_of_client(session, bank_a.id)),
                          asof=date(2025, 1, 1))
    assert apres.confidence == "insufficient_history"
