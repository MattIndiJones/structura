"""
Generate AMC Diagnostic Excel workbooks from LUKB raw exports.

Usage:
    python generate_amc_excel.py

Reads:  C:/Users/phili/Downloads/UTI/Data/{ISIN}/
Writes: C:/Users/phili/Downloads/UTI/Data/{ISIN}/{ISIN}_AMC_Diagnostic_Master.xlsx

Skips ISINs that already have an Excel file.
"""
from __future__ import annotations

import json
import math
import os
import re
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(r"C:\Users\phili\Downloads\UTI\Data")

# ISINs to process (skip CH1473733959 — already has the reference file)
TARGET_ISINS = [
    "CH1352587708",  # New Financials CHF
    "CH1352587716",  # New Financials EUR
    "CH1352587724",  # New Financials USD  — already has CORRECTED
    "CH1473731680",  # NEU Life USD        — already has CORRECTED
    "CH1473733934",  # NEU Infrastructures CHF — already has file
    "CH1473736143",  # NEU Infrastructures H USD — MISSING
]

# ── Styles ─────────────────────────────────────────────────────────────
DARK_BG  = "1E293B"
BLUE_HDR = "1E40AF"
GREEN    = "065F46"
AMBER    = "92400E"
GRAY     = "374151"

def _hdr_fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, color="FFFFFF", size=9):
    return Font(bold=bold, color=color, size=size, name="Calibri")

def _border():
    s = Side(style="thin", color="4B5563")
    return Border(left=s, right=s, top=s, bottom=s)

def _write_header_row(ws, row: int, values: list, fill_color: str = BLUE_HDR):
    for col, val in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=val)
        c.fill = _hdr_fill(fill_color)
        c.font = _font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _border()

def _write_row(ws, row: int, values: list, bold=False, fill: str = None):
    for col, val in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=val)
        c.font = _font(bold=bold, color="E2E8F0")
        c.alignment = Alignment(vertical="center")
        c.border = _border()
        if fill:
            c.fill = _hdr_fill(fill)

def _set_col_widths(ws, widths: list):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── Data loading ───────────────────────────────────────────────────────

def _load_def(folder: Path) -> dict:
    isin = folder.name
    with open(folder / f"{isin} Def.txt", encoding="utf-8") as f:
        d = json.load(f)
    return d["data"]["products"]["items"][0]


def _load_nav(folder: Path) -> pd.DataFrame:
    csv_files = list(folder.glob("*.csv"))
    if not csv_files:
        return pd.DataFrame()
    df = pd.read_csv(csv_files[0], encoding="utf-8")
    df.columns = [c.strip() for c in df.columns]
    date_col  = next((c for c in df.columns if "date" in c.lower()), df.columns[0])
    price_col = next((c for c in df.columns if "price" in c.lower() or "nav" in c.lower()), df.columns[1])
    out_col   = next((c for c in df.columns if "outstanding" in c.lower() or "quantity" in c.lower()), None)

    df = df.rename(columns={date_col: "Date", price_col: "NAV"})
    if out_col:
        df = df.rename(columns={out_col: "Outstanding Quantity"})
    else:
        df["Outstanding Quantity"] = None

    # Parse dates — format can be DD.MM.YYYY or YYYY-MM-DD
    df["Date"] = pd.to_datetime(df["Date"], dayfirst=True, errors="coerce")
    df = df.dropna(subset=["Date"]).sort_values("Date")
    df["NAV"] = pd.to_numeric(df["NAV"], errors="coerce")
    df = df.dropna(subset=["NAV"])

    # Derived columns
    df["Daily Return"]      = df["NAV"].pct_change()
    df["Cumulative Return"] = df["NAV"] / df["NAV"].iloc[0] - 1
    df["Running Peak"]      = df["NAV"].cummax()
    df["Drawdown"]          = df["NAV"] / df["Running Peak"] - 1
    df["Month"]             = df["Date"].dt.to_period("M").astype(str)
    return df


def _load_orders(folder: Path) -> list[dict]:
    merged_files = list(folder.glob("merged-*.json"))
    if not merged_files:
        return []
    with open(merged_files[0], encoding="utf-8") as f:
        d = json.load(f)
    items = d["data"]["orders"]["items"]
    orders = []
    for o in items:
        if o.get("state") not in ("Done", "Discarded"):
            continue
        ep = o.get("executionPrice") or {}
        price_val = float(ep.get("amount") or ep.get("amountDirty") or 0)
        price_dirty = float(ep.get("amountDirty") or ep.get("amount") or 0)
        price_ccy  = ep.get("currency", "")
        price_type = ep.get("priceType", "")
        fx = float(o.get("usedFxRate") or 1)
        oq = float(o.get("orderedQuantity") or 0)
        eq = float(o.get("executedQuantity") or 0)
        und = o.get("underlying") or {}
        # parse datetime
        dt_str = o.get("creationDateTime", "")
        try:
            dt = datetime.fromisoformat(dt_str.replace("Z", "+00:00"))
        except Exception:
            dt = None

        local_notional = eq * price_val
        prod_notional  = local_notional * fx  # product-ccy notional (signed)

        orders.append({
            "ID":                       o.get("id", ""),
            "State":                    o.get("state", ""),
            "DateTime":                 dt_str,
            "Date":                     dt.strftime("%Y-%m-%d") if dt else "",
            "Side":                     "BUY" if oq > 0 else "SELL",
            "Underlying Name":          und.get("name", ""),
            "Underlying ISIN":          und.get("isin", ""),
            "Underlying Currency":      und.get("currency", ""),
            "Ordered Qty":              oq,
            "Executed Qty":             eq,
            "Execution Price":          price_val,
            "Price Dirty":              price_dirty,
            "Trade Currency":           price_ccy,
            "Price Type":               price_type,
            "FX Rate":                  fx,
            "Local Notional Signed":    round(local_notional, 6),
            "Prod Notional Signed":     round(prod_notional, 6),
            "Prod Notional Abs":        round(abs(prod_notional), 6),
        })
    return orders


def _trade_summary(orders: list[dict]) -> pd.DataFrame:
    """Aggregate per underlying."""
    done = [o for o in orders if o["State"] == "Done" and o["Executed Qty"]]
    by_isin: dict[str, dict] = {}
    for o in done:
        isin = o["Underlying ISIN"]
        if isin not in by_isin:
            by_isin[isin] = {
                "Underlying Name": o["Underlying Name"],
                "ISIN": isin,
                "Currency": o["Underlying Currency"],
                "Trades": 0,
                "Buy Qty": 0.0, "Sell Qty": 0.0, "Net Qty": 0.0,
                "Buy Notional": 0.0, "Sell Notional": 0.0,
            }
        r = by_isin[isin]
        r["Trades"] += 1
        if o["Executed Qty"] > 0:
            r["Buy Qty"] += o["Executed Qty"]
            r["Buy Notional"] += o["Prod Notional Abs"]
        else:
            r["Sell Qty"] += o["Executed Qty"]
            r["Sell Notional"] += o["Prod Notional Abs"]
        r["Net Qty"] += o["Executed Qty"]

    rows = []
    for r in by_isin.values():
        rows.append({
            "Underlying Name": r["Underlying Name"],
            "ISIN": r["ISIN"],
            "Currency": r["Currency"],
            "Trades": r["Trades"],
            "Buy Qty": round(r["Buy Qty"], 4),
            "Sell Qty": round(r["Sell Qty"], 4),
            "Net Qty": round(r["Net Qty"], 4),
            "Buy Notional": round(r["Buy Notional"], 2),
            "Sell Notional": round(-r["Sell Notional"], 2),
            "Gross Traded": round(r["Buy Notional"] + r["Sell Notional"], 2),
        })
    rows.sort(key=lambda r: r["Gross Traded"], reverse=True)
    return pd.DataFrame(rows)


def _composition_df(product: dict) -> pd.DataFrame:
    """Build Current_Composition_API sheet content."""
    nav_val = float(product["netAssetValue"]["value"])
    outstanding = float(product.get("outstandingQuantity") or 0)
    total_aum = nav_val * outstanding
    rows = []
    for c in product.get("components", []):
        und = c.get("underlying") or {}
        weight_raw = float(c.get("weight") or 0)
        position  = float(c.get("position") or 0)
        # We don't have individual prices from Def.txt — use NAV as proxy
        # position_x_nav ≈ weight * total_aum
        pos_x_nav = weight_raw * total_aum
        rows.append({
            "Underlying Name":         und.get("name", ""),
            "ISIN":                    und.get("isin", ""),
            "Currency":                und.get("currency", ""),
            "Position":                round(position, 6),
            "API Weight Raw":          round(weight_raw, 6),
            "API Weight %":            round(weight_raw, 6),
            "PRICE":                   round(nav_val, 5),
            "Position x Nav":          round(pos_x_nav, 2),
            "Weight Per Certificat":   round(weight_raw / outstanding if outstanding else 0, 6),
        })
    rows.sort(key=lambda r: r["API Weight %"], reverse=True)
    return pd.DataFrame(rows)


# ── Sheet builders ──────────────────────────────────────────────────────

def _sheet_dashboard(wb: Workbook, product: dict, nav_df: pd.DataFrame,
                     orders: list[dict], isin: str):
    ws = wb.create_sheet("Dashboard")
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for c in row:
            c.fill = _hdr_fill(DARK_BG)

    nav_date = product["netAssetValue"]["date"][:10]
    nav_val  = float(product["netAssetValue"]["value"])
    n_nav    = len(nav_df)
    n_trades = len([o for o in orders if o["State"] == "Done"])
    n_unique = len(set(o["Underlying ISIN"] for o in orders if o["State"] == "Done"))
    n_active_days = len(set(o["Date"] for o in orders if o["State"] == "Done"))
    n_buy  = len([o for o in orders if o["State"] == "Done" and o["Executed Qty"] > 0])
    n_sell = len([o for o in orders if o["State"] == "Done" and o["Executed Qty"] < 0])

    # Title
    ws.merge_cells("A1:E1")
    c = ws["A1"]
    c.value = f"AMC Diagnostic Workbook - {isin}"
    c.font = Font(bold=True, size=13, color="FFFFFF", name="Calibri")
    c.fill = _hdr_fill(BLUE_HDR)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    meta = [
        ("Product name",  product.get("name", "")),
        ("ISIN",          isin),
        ("Currency",      product.get("currency", "")),
        ("Issue date",    product.get("issueDate", "")[:10] if product.get("issueDate") else ""),
        ("NAV start",     str(nav_df["Date"].iloc[0].date()) if n_nav else ""),
        ("NAV end",       str(nav_df["Date"].iloc[-1].date()) if n_nav else ""),
        ("Outstanding",   product.get("outstandingQuantity", "")),
        ("NAV as of",     nav_date),
        ("NAV value",     round(nav_val, 4)),
    ]
    status = [
        ("Data status",            "Loaded"),
        ("NAV observations",       n_nav),
        ("Trade count",            n_trades),
        ("Unique traded securities", n_unique),
        ("Active trade days",      n_active_days),
        ("Buy / Sell trades",      f"{n_buy} / {n_sell}"),
    ]

    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True, color="94A3B8", name="Calibri", size=9)
        ws.cell(row=i, column=2, value=v).font  = Font(color="E2E8F0", name="Calibri", size=9)

    for i, (k, v) in enumerate(status, start=3):
        ws.cell(row=i, column=4, value=k).font = Font(bold=True, color="94A3B8", name="Calibri", size=9)
        ws.cell(row=i, column=5, value=v).font = Font(color="E2E8F0", name="Calibri", size=9)

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 20


def _sheet_nav(wb: Workbook, nav_df: pd.DataFrame):
    ws = wb.create_sheet("NAV_History")
    ws.sheet_view.showGridLines = False

    headers = ["Date", "NAV", "Outstanding Quantity", "Daily Return",
               "Cumulative Return", "Running Peak", "Drawdown", "Month"]
    _write_header_row(ws, 1, headers)
    _set_col_widths(ws, [14, 14, 20, 14, 18, 14, 14, 10])

    def _safe(v):
        if v is None or (isinstance(v, float) and math.isnan(v)):
            return None
        return v

    for i, row in nav_df.iterrows():
        r = ws.max_row + 1
        vals = [
            row["Date"].date() if pd.notna(row["Date"]) else None,
            round(float(row["NAV"]), 6) if pd.notna(row["NAV"]) else None,
            round(float(row["Outstanding Quantity"]), 0) if pd.notna(row.get("Outstanding Quantity")) else None,
            _safe(round(float(row["Daily Return"]), 6)) if pd.notna(row.get("Daily Return")) else None,
            _safe(round(float(row["Cumulative Return"]), 6)) if pd.notna(row.get("Cumulative Return")) else None,
            _safe(round(float(row["Running Peak"]), 6)) if pd.notna(row.get("Running Peak")) else None,
            _safe(round(float(row["Drawdown"]), 6)) if pd.notna(row.get("Drawdown")) else None,
            str(row["Month"]) if "Month" in row else None,
        ]
        _write_row(ws, r, vals)


def _sheet_transactions(wb: Workbook, orders: list[dict], prod_ccy: str):
    ws = wb.create_sheet("Transactions")
    ws.sheet_view.showGridLines = False

    fx_col = f"FX Rate to {prod_ccy}"
    local_col = "Local Notional Signed"
    prod_col_s = f"{prod_ccy} Notional Signed"
    prod_col_a = f"{prod_ccy} Notional Abs"

    headers = ["ID", "State", "DateTime", "Date", "Side",
               "Underlying Name", "Underlying ISIN", "Underlying Currency",
               "Ordered Qty", "Executed Qty", "Execution Price", "Price Dirty",
               "Trade Currency", "Price Type",
               fx_col, local_col, prod_col_s, prod_col_a]
    _write_header_row(ws, 1, headers)
    _set_col_widths(ws, [38, 10, 26, 12, 6, 30, 14, 10, 10, 10, 15, 12, 10, 10, 12, 18, 18, 18])

    done_orders = [o for o in orders if o["State"] == "Done"]
    for o in done_orders:
        r = ws.max_row + 1
        _write_row(ws, r, [
            o["ID"], o["State"], o["DateTime"], o["Date"], o["Side"],
            o["Underlying Name"], o["Underlying ISIN"], o["Underlying Currency"],
            o["Ordered Qty"], o["Executed Qty"],
            o["Execution Price"], o["Price Dirty"],
            o["Trade Currency"], o["Price Type"],
            o["FX Rate"],
            o["Local Notional Signed"], o["Prod Notional Signed"], o["Prod Notional Abs"],
        ])


def _sheet_trade_summary(wb: Workbook, orders: list[dict], prod_ccy: str):
    ws = wb.create_sheet("Trade_Summary")
    ws.sheet_view.showGridLines = False

    buy_col  = f"Buy Notional {prod_ccy}"
    sell_col = f"Sell Notional {prod_ccy}"
    gross_col = f"Gross Traded {prod_ccy}"

    df = _trade_summary(orders)
    if df.empty:
        ws.cell(1, 1, "No trades")
        return

    headers = ["Underlying Name", "ISIN", "Currency", "Trades",
               "Buy Qty", "Sell Qty", "Net Qty",
               buy_col, sell_col, gross_col]
    _write_header_row(ws, 1, headers)
    _set_col_widths(ws, [35, 14, 10, 8, 10, 10, 10, 18, 18, 18])

    for _, row in df.iterrows():
        r = ws.max_row + 1
        _write_row(ws, r, [
            row["Underlying Name"], row["ISIN"], row["Currency"], row["Trades"],
            row["Buy Qty"], row["Sell Qty"], row["Net Qty"],
            row["Buy Notional"], row["Sell Notional"], row["Gross Traded"],
        ])


def _sheet_composition(wb: Workbook, product: dict):
    ws = wb.create_sheet("Current_Composition_API")
    ws.sheet_view.showGridLines = False

    headers = ["Underlying Name", "ISIN", "Currency", "Position",
               "API Weight Raw", "API Weight %", "PRICE",
               "Position x Nav", "Weight Per Certificat", "Note"]
    _write_header_row(ws, 1, headers)
    _set_col_widths(ws, [35, 14, 10, 14, 14, 14, 12, 16, 20, 55])

    df = _composition_df(product)
    for _, row in df.iterrows():
        r = ws.max_row + 1
        _write_row(ws, r, [
            row["Underlying Name"], row["ISIN"], row["Currency"],
            row["Position"], row["API Weight Raw"], row["API Weight %"],
            row["PRICE"], row["Position x Nav"], row["Weight Per Certificat"],
            None,
        ])


def _sheet_factsheet_summary(wb: Workbook, product: dict, nav_df: pd.DataFrame):
    ws = wb.create_sheet("Factsheet_Summary")
    ws.sheet_view.showGridLines = False
    nav_date = product["netAssetValue"]["date"][:10]
    nav_val  = float(product["netAssetValue"]["value"])
    issue = (product.get("issueDate") or "")[:10]

    rows = [
        (f"Factsheet Extract - data as of {nav_date}", ""),
        ("", ""),
        ("Product",     product.get("name", "")),
        ("ISIN",        product.get("isin", "")),
        ("Currency",    product.get("currency", "")),
        ("NAV Date",    nav_date),
        ("NAV",         round(nav_val, 4)),
        ("Issue Date",  issue),
        ("Outstanding", product.get("outstandingQuantity", "")),
        ("", ""),
        ("NAV observations", len(nav_df)),
        ("NAV start",   str(nav_df["Date"].iloc[0].date()) if len(nav_df) else ""),
        ("NAV end",     str(nav_df["Date"].iloc[-1].date()) if len(nav_df) else ""),
    ]
    for i, (k, v) in enumerate(rows, start=1):
        ws.cell(i, 1, k).font = Font(bold=bool(k), color="94A3B8", name="Calibri", size=9)
        ws.cell(i, 2, v).font = Font(color="E2E8F0", name="Calibri", size=9)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 65


def _sheet_factsheet_exposures(wb: Workbook, product: dict):
    ws = wb.create_sheet("Factsheet_Exposures")
    ws.sheet_view.showGridLines = False

    # Derive currency breakdown from composition
    ccy_weights: dict[str, float] = {}
    for c in product.get("components", []):
        ccy = (c.get("underlying") or {}).get("currency", "?")
        ccy_weights[ccy] = ccy_weights.get(ccy, 0.0) + float(c.get("weight") or 0)

    # We don't have sector/country data — mark as pending factsheet
    _write_header_row(ws, 1, ["Sector Exposure", "", "", "Country Exposure", "", "", "Currency Exposure", ""], GRAY)
    _write_row(ws, 2, ["Sector", "Weight", "", "Country", "Weight", "", "Currency", "Weight"], bold=True)

    ws.cell(3, 1, "Pending factsheet upload").font = Font(italic=True, color="64748B", name="Calibri", size=9)
    ws.cell(3, 4, "Pending factsheet upload").font = Font(italic=True, color="64748B", name="Calibri", size=9)

    for i, (ccy, w) in enumerate(sorted(ccy_weights.items(), key=lambda x: -x[1]), start=3):
        ws.cell(i, 7, ccy).font = Font(color="E2E8F0", name="Calibri", size=9)
        ws.cell(i, 8, round(w, 4)).font = Font(color="E2E8F0", name="Calibri", size=9)

    _set_col_widths(ws, [28, 10, 4, 28, 10, 4, 14, 10])


def _sheet_factsheet_toplists(wb: Workbook, product: dict):
    ws = wb.create_sheet("Factsheet_TopLists")
    ws.sheet_view.showGridLines = False

    comps = sorted(product.get("components", []),
                   key=lambda c: float(c.get("weight") or 0), reverse=True)

    _write_header_row(ws, 1, ["Top Holdings (by weight)", "", "", "", ""])
    _write_row(ws, 2, ["Name", "ISIN", "Currency", "Weight", ""], bold=True)
    for i, c in enumerate(comps[:10], start=3):
        und = c.get("underlying") or {}
        _write_row(ws, i, [und.get("name"), und.get("isin"), und.get("currency"),
                            round(float(c.get("weight") or 0), 4), ""])
    _set_col_widths(ws, [35, 14, 10, 10, 10])


def _sheet_data_inventory(wb: Workbook, product: dict, nav_df: pd.DataFrame,
                          orders: list[dict]):
    ws = wb.create_sheet("Data_Inventory")
    ws.sheet_view.showGridLines = False

    nav_count = len(nav_df)
    nav_start = str(nav_df["Date"].iloc[0].date()) if nav_count else "—"
    nav_end   = str(nav_df["Date"].iloc[-1].date()) if nav_count else "—"
    order_count = len([o for o in orders if o["State"] == "Done"])

    headers = ["Dataset", "Source", "Status", "Comments"]
    _write_header_row(ws, 1, headers)
    _set_col_widths(ws, [25, 35, 10, 60])

    rows = [
        ("NAV time series",      "CSV export uploaded by user",            "Loaded",
         f"{nav_count} observations from {nav_start} to {nav_end}"),
        ("Orders / transactions", "GraphQL JSON export (merged file)",      "Loaded",
         f"{order_count} orders loaded from merged export"),
        ("Current composition",   "Def.txt GraphQL snapshot",              "Loaded",
         f"{len(product.get('components',[]))} positions as of {product['netAssetValue']['date'][:10]}"),
        ("Factsheet summary",     "Def.txt",                               "Partial",
         "Sector/country exposures pending — currency breakdown derived from composition"),
        ("Benchmark",             "Client / advisor to confirm",           "Missing",
         "Required for relative performance and information ratio"),
        ("Fama-French factors",   "Kenneth French Data Library",           "Missing",
         "Required for FF5 + Momentum regression — import via 'Mettre à jour les facteurs'"),
        ("Security price histories", "Market data provider",               "Optional",
         "Required for precise position-level attribution"),
    ]
    for row_data in rows:
        r = ws.max_row + 1
        status_fill = GREEN if row_data[2] == "Loaded" else AMBER if row_data[2] == "Partial" else GRAY
        _write_row(ws, r, list(row_data))
        ws.cell(r, 3).fill = _hdr_fill(status_fill)


def _sheet_methodology(wb: Workbook):
    ws = wb.create_sheet("Methodology")
    ws.sheet_view.showGridLines = False

    headers = ["Topic", "Method / Formula", "Comment"]
    _write_header_row(ws, 1, headers)
    _set_col_widths(ws, [28, 50, 40])

    rows = [
        ("NAV performance",       "Daily return = NAV_t / NAV_(t-1) - 1",        "Implemented in NAV_History"),
        ("Cumulative return",     "NAV_t / NAV_start - 1",                        "Implemented in NAV_History"),
        ("Drawdown",              "NAV_t / running peak - 1",                     "Implemented in NAV_History"),
        ("Trade notional",        "Executed qty × execution price × FX rate",     "Signed product-ccy notional in Transactions"),
        ("Turnover proxy",        "Gross traded / approximate AUM",               "Can be refined once average AUM is stabilized"),
        ("Factor regression",     "AMC excess return = alpha + beta × factors + ε", "Pending benchmark/factor dataset"),
        ("Attribution limitation","Orders alone do not provide daily M2M P&L",   "Need security price history for exact contribution"),
        ("FIFO P&L",              "FIFO matching on executed orders",             "Implemented in study engine (Bloc A→D)"),
    ]
    for row_data in rows:
        r = ws.max_row + 1
        _write_row(ws, r, list(row_data))


def _sheet_checks(wb: Workbook, nav_df: pd.DataFrame, orders: list[dict]):
    ws = wb.create_sheet("Checks")
    ws.sheet_view.showGridLines = False

    headers = ["Check", "Value", "Status"]
    _write_header_row(ws, 1, headers)
    _set_col_widths(ws, [30, 20, 10])

    ids = [o["ID"] for o in orders if o["State"] == "Done"]
    dup = len(ids) - len(set(ids))
    nav_start = nav_df["Date"].iloc[0].date() if len(nav_df) else None
    nav_end   = nav_df["Date"].iloc[-1].date() if len(nav_df) else None
    order_dates = [o["Date"] for o in orders if o["State"] == "Done" and o["Date"]]

    checks = [
        ("NAV rows loaded",           len(nav_df),          "OK" if len(nav_df) > 0 else "WARN"),
        ("Transaction rows loaded",   len([o for o in orders if o["State"] == "Done"]),
                                                             "OK" if orders else "WARN"),
        ("Duplicate transaction IDs", dup,                   "OK" if dup == 0 else "WARN"),
        ("First NAV date",            str(nav_start) if nav_start else "—", "OK" if nav_start else "WARN"),
        ("Last NAV date",             str(nav_end) if nav_end else "—",     "OK" if nav_end else "WARN"),
        ("First transaction date",    min(order_dates) if order_dates else "—",
                                                             "OK" if order_dates else "WARN"),
        ("Last transaction date",     max(order_dates) if order_dates else "—",
                                                             "OK" if order_dates else "WARN"),
    ]
    for row_data in checks:
        r = ws.max_row + 1
        fill = GREEN if row_data[2] == "OK" else AMBER
        _write_row(ws, r, list(row_data))
        ws.cell(r, 3).fill = _hdr_fill(fill)


# ── Main builder ────────────────────────────────────────────────────────

def build_excel(isin: str) -> Path:
    folder = BASE_DIR / isin
    product  = _load_def(folder)
    nav_df   = _load_nav(folder)
    orders   = _load_orders(folder)
    prod_ccy = product.get("currency", "USD")

    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    _sheet_dashboard(wb, product, nav_df, orders, isin)
    _sheet_nav(wb, nav_df)
    _sheet_transactions(wb, orders, prod_ccy)
    _sheet_trade_summary(wb, orders, prod_ccy)
    _sheet_composition(wb, product)
    _sheet_factsheet_summary(wb, product, nav_df)
    _sheet_factsheet_exposures(wb, product)
    _sheet_factsheet_toplists(wb, product)
    _sheet_data_inventory(wb, product, nav_df, orders)
    _sheet_methodology(wb)
    _sheet_checks(wb, nav_df, orders)

    out_path = folder / f"{isin}_AMC_Diagnostic_Master.xlsx"
    wb.save(str(out_path))
    return out_path


def main():
    import sys
    sys.stdout.reconfigure(encoding="utf-8")
    for isin in TARGET_ISINS:
        folder = BASE_DIR / isin
        existing = list(folder.glob(f"{isin}*.xlsx"))
        if existing:
            print(f"[SKIP] {isin} - Excel exists: {existing[0].name}")
            continue
        print(f"[BUILD] {isin} ...", end="", flush=True)
        try:
            path = build_excel(isin)
            print(f" done: {path.name}")
        except Exception as e:
            print(f" ERROR: {e}")


if __name__ == "__main__":
    main()
